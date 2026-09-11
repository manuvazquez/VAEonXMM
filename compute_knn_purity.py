from __future__ import annotations

"""Compute kNN purity summaries for Gaussian latent representations.

Each row in a latent.parquet file is treated as one Gaussian distribution:
`z` is the mean vector and `std` is the diagonal standard deviation vector.
The script keeps only manually labeled rows, computes a distance matrix under
one or more Gaussian distances, and writes one Excel summary per distance.
"""

import argparse
import os
import re
from pathlib import Path

_cache_dir = Path(".cache")
_cache_dir.mkdir(parents=True, exist_ok=True)
os.environ.setdefault("MPLCONFIGDIR", str(_cache_dir / "matplotlib"))
os.environ.setdefault("XDG_CACHE_HOME", str(_cache_dir))

import matplotlib.pyplot as plt
import numpy as np
import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import PatternFill

NEGATIVE_LABEL = "NOT interesting"
POSITIVE_LABEL = "Interesting"
PARAMETER_PATTERN = re.compile(
    r"beta_(?P<beta>[^_]+)_dim_z_(?P<dim_z>\d+)"
)
DIM_Z_PATTERN = re.compile(r"(?:^|_)z_(?P<dim_z>\d+)(?:_|$)")
PURITY_COLUMNS = (
    "avg_purity_positive",
    "avg_purity_negative",
    "avg_purity",
)


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description=(
            "Compute class-wise kNN purity for labeled samples in every "
            "latent.parquet file under a root folder."
        )
    )
    parser.add_argument(
        "--root",
        type=Path,
        default=Path("datasets/rebuttal"),
        help="Root folder containing latent.parquet or vae.parquet files.",
    )
    parser.add_argument(
        "--annotations",
        type=Path,
        default=Path("annotations/annotated_windows.parquet"),
        help="Annotations file used when model outputs do not contain labels.",
    )
    parser.add_argument(
        "--output-dir",
        type=Path,
        default=Path("knn_purity"),
        help="Folder where the summary Excel files will be written.",
    )
    parser.add_argument(
        "-k",
        type=int,
        default=10,
        help="Number of nearest neighbours used for purity.",
    )
    parser.add_argument(
        "--distances",
        nargs="+",
        default=[
            "wasserstein",
            "euclidean_means",
            "mahalanobis_pooled",
            "jensen_shannon_mc",
            "bhattacharyya",
            "symmetric_kl",
        ],
        help="Distance names to compute.",
    )
    parser.add_argument(
        "--js-samples",
        type=int,
        default=64,
        help="Monte Carlo samples per distribution for jensen_shannon_mc.",
    )
    parser.add_argument(
        "--seed",
        type=int,
        default=0,
        help="Random seed for Monte Carlo distances.",
    )
    return parser.parse_args()


def stack_vector_column(series: pd.Series, column_name: str) -> np.ndarray:
    """Turn a pandas column of vectors into a 2D numpy array."""
    lengths = series.map(len)
    if lengths.nunique() != 1:
        raise ValueError(
            f"Column {column_name!r} has mixed vector lengths: "
            f"{lengths.value_counts().sort_index().to_dict()}"
        )
    return np.vstack(series.to_numpy()).astype(float, copy=False)


def parse_parameters(folder_name: str) -> tuple[float, int]:
    """Extract beta and dim_z values from a run folder name."""
    match = PARAMETER_PATTERN.search(folder_name)
    if match is not None:
        return float(match.group("beta")), int(match.group("dim_z"))

    dim_z_match = DIM_Z_PATTERN.search(folder_name)
    if dim_z_match is not None:
        return np.nan, int(dim_z_match.group("dim_z"))

    return np.nan, -1


def component_columns(columns: pd.Index, prefix: str) -> list[str]:
    """Return numbered component columns sorted by component index."""
    pattern = re.compile(rf"^{re.escape(prefix)}_(\d+)$")
    matches = [
        (int(match.group(1)), column)
        for column in columns
        if (match := pattern.match(column))
    ]
    return [column for _, column in sorted(matches)]


def data_paths(root: Path) -> list[Path]:
    """Find supported model output files, preferring latent over vae files."""
    paths_by_folder = {}
    for path in sorted(root.glob("**/vae.parquet")):
        paths_by_folder[path.parent] = path
    for path in sorted(root.glob("**/latent.parquet")):
        run_folder = (
            path.parent.parent if path.parent.name == "corpus" else path.parent
        )
        paths_by_folder[run_folder] = path
    return [paths_by_folder[key] for key in sorted(paths_by_folder)]


def run_folder(path: Path) -> Path:
    """Return the experiment folder for either supported file layout."""
    if path.name == "latent.parquet" and path.parent.name == "corpus":
        return path.parent.parent
    return path.parent


def load_labeled_samples(path: Path, annotations_path: Path) -> pd.DataFrame:
    """Load labels and latent Gaussian parameters from a model output file."""
    df = pd.read_parquet(path)
    if {"Label", "z", "std"}.issubset(df.columns):
        return df

    z_columns = component_columns(df.columns, "z")
    std_columns = component_columns(df.columns, "z_sd")
    if not z_columns or len(z_columns) != len(std_columns):
        raise ValueError(f"Could not find matching z/z_sd columns in {path}")

    if "Label" not in df.columns:
        annotation_columns = [
            "start",
            "end",
            "tar file",
            "source",
            "signal",
            "Label",
        ]
        annotations = pd.read_parquet(annotations_path,
                                      columns=annotation_columns)
        merge_columns = [column for column in annotation_columns[:-1]
                         if column in df.columns]
        df = df.merge(annotations, on=merge_columns, how="left",
                      validate="many_to_one")

    df = df.copy()
    df["z"] = list(df[z_columns].to_numpy(dtype=float))
    df["std"] = list(df[std_columns].to_numpy(dtype=float))
    return df


def diagonal_gaussian_w2_distance_matrix(
    means: np.ndarray, stds: np.ndarray
) -> np.ndarray:
    """2-Wasserstein distances for diagonal Gaussian distributions."""
    # For diagonal Gaussians, W2^2 is just the squared Euclidean distance
    # between means plus the squared Euclidean distance between std vectors.
    mu_sq = np.sum((means[:, None, :] - means[None, :, :]) ** 2, axis=2)
    std_sq = np.sum((stds[:, None, :] - stds[None, :, :]) ** 2, axis=2)
    return np.sqrt(mu_sq + std_sq)


def euclidean_means_distance_matrix(
    means: np.ndarray, stds: np.ndarray
) -> np.ndarray:
    """Plain Euclidean distance between mean vectors; ignores uncertainty."""
    del stds
    return np.sqrt(
        np.sum((means[:, None, :] - means[None, :, :]) ** 2, axis=2)
    )


def mahalanobis_pooled_distance_matrix(
    means: np.ndarray, stds: np.ndarray
) -> np.ndarray:
    """Mahalanobis distance with the average covariance of each pair."""
    variances = stds**2
    pooled_variances = 0.5 * (variances[:, None, :] + variances[None, :, :])
    mean_sq = (means[:, None, :] - means[None, :, :]) ** 2
    return np.sqrt(np.sum(mean_sq / pooled_variances, axis=2))


def bhattacharyya_distance_matrix(
    means: np.ndarray, stds: np.ndarray
) -> np.ndarray:
    """Closed-form Bhattacharyya distance for diagonal Gaussians."""
    variances = stds**2
    avg_variances = 0.5 * (variances[:, None, :] + variances[None, :, :])
    mean_sq = (means[:, None, :] - means[None, :, :]) ** 2
    mean_term = 0.125 * np.sum(mean_sq / avg_variances, axis=2)
    var_term = 0.5 * np.sum(
        np.log(
            avg_variances
            / np.sqrt(variances[:, None, :] * variances[None, :, :])
        ),
        axis=2,
    )
    return mean_term + var_term


def symmetric_kl_distance_matrix(
    means: np.ndarray, stds: np.ndarray
) -> np.ndarray:
    """Symmetric KL divergence: 0.5 * (KL(p || q) + KL(q || p))."""
    variances = stds**2
    mean_sq = (means[:, None, :] - means[None, :, :]) ** 2
    kl_ij = 0.5 * np.sum(
        np.log(variances[None, :, :] / variances[:, None, :])
        + (variances[:, None, :] + mean_sq) / variances[None, :, :]
        - 1.0,
        axis=2,
    )
    return 0.5 * (kl_ij + kl_ij.T)


def diagonal_gaussian_logpdf(
    samples: np.ndarray, means: np.ndarray, stds: np.ndarray
) -> np.ndarray:
    """Evaluate many diagonal Gaussian log densities at many sample points."""
    variances = stds**2
    centered_sq = (samples[:, None, :] - means[None, :, :]) ** 2
    return -0.5 * np.sum(
        centered_sq / variances[None, :, :]
        + np.log(2.0 * np.pi * variances[None, :, :]),
        axis=2,
    )


def jensen_shannon_mc_distance_matrix(
    means: np.ndarray,
    stds: np.ndarray,
    samples_per_distribution: int,
    rng: np.random.Generator,
) -> np.ndarray:
    """Monte Carlo estimate of Jensen-Shannon distance.

    Exact Jensen-Shannon distance between two continuous Gaussian densities
    requires the entropy of a Gaussian mixture, which has no simple closed
    form. We estimate E_p[log p - log m] and E_q[log q - log m] by sampling,
    where m = 0.5 * (p + q), then return sqrt(JS divergence).
    """
    if samples_per_distribution < 1:
        raise ValueError("--js-samples must be at least 1")

    n_samples = len(means)
    expectations = np.empty((n_samples, n_samples), dtype=float)
    for i in range(n_samples):
        # Samples are drawn from distribution i and compared against every
        # distribution j, giving one row of E_i[log p_i - log m_ij].
        samples = rng.normal(
            loc=means[i],
            scale=stds[i],
            size=(samples_per_distribution, means.shape[1]),
        )
        log_probs = diagonal_gaussian_logpdf(samples, means, stds)
        log_p = log_probs[:, [i]]
        log_m = np.logaddexp(log_p, log_probs) - np.log(2.0)
        expectations[i] = np.mean(log_p - log_m, axis=0)

    js = 0.5 * (expectations + expectations.T)
    np.fill_diagonal(js, 0.0)
    return np.sqrt(np.maximum(0.0, js))


def distance_matrix(
    name: str,
    means: np.ndarray,
    stds: np.ndarray,
    js_samples: int,
    rng: np.random.Generator,
) -> np.ndarray:
    """Dispatch a human-readable distance name to its implementation."""
    distance_functions = {
        "wasserstein": diagonal_gaussian_w2_distance_matrix,
        "euclidean_means": euclidean_means_distance_matrix,
        "mahalanobis_pooled": mahalanobis_pooled_distance_matrix,
        "bhattacharyya": bhattacharyya_distance_matrix,
        "symmetric_kl": symmetric_kl_distance_matrix,
    }
    if name == "jensen_shannon_mc":
        return jensen_shannon_mc_distance_matrix(means, stds, js_samples, rng)
    if name not in distance_functions:
        known_names = sorted([*distance_functions, "jensen_shannon_mc"])
        raise ValueError(
            f"Unknown distance {name!r}. Known distances: {known_names}"
        )
    return distance_functions[name](means, stds)


def knn_purity(
    labels: np.ndarray, distances: np.ndarray, k: int
) -> np.ndarray:
    """Compute the fraction of same-label k nearest neighbours."""
    n_samples = len(labels)
    actual_k = min(k, n_samples - 1)
    if actual_k <= 0:
        return np.full(n_samples, np.nan)

    # The first neighbour after sorting is the sample itself at distance zero,
    # so we skip column 0 and keep the next k positions.
    neighbour_order = np.argsort(distances, axis=1)[:, 1:actual_k + 1]
    neighbour_labels = labels[neighbour_order]
    return (neighbour_labels == labels[:, None]).mean(axis=1)


def summarize_file(
    path: Path,
    root: Path,
    annotations_path: Path,
    k: int,
    distance_name: str,
    js_samples: int,
    rng: np.random.Generator,
) -> dict[str, object]:
    """Compute all summary numbers for one model output file and distance."""
    df = load_labeled_samples(path, annotations_path)

    # Keep only rows with the two labels used in the binary purity analysis.
    df = df[df["Label"].isin([NEGATIVE_LABEL, POSITIVE_LABEL])].copy()

    label_map = {NEGATIVE_LABEL: 0, POSITIVE_LABEL: 1}
    labels = df["Label"].map(label_map).to_numpy()
    means = stack_vector_column(df["z"], "z")
    stds = stack_vector_column(df["std"], "std")
    if distance_name != "euclidean_means" and stds.shape[1] != means.shape[1]:
        raise ValueError(
            f"{distance_name} requires std vectors with dimension "
            f"{means.shape[1]}, found {stds.shape[1]}"
        )

    distances = distance_matrix(distance_name, means, stds, js_samples, rng)
    purities = knn_purity(labels, distances, k)

    positive_mask = labels == 1
    negative_mask = labels == 0
    folder_name = run_folder(path).relative_to(root).as_posix()
    beta, dim_z = parse_parameters(folder_name)

    # Class-wise averages are useful when the classes are imbalanced. The
    # global average below is the sample-weighted average over all rows.
    avg_purity_positive = (
        float(np.nanmean(purities[positive_mask]))
        if positive_mask.any()
        else np.nan
    )
    avg_purity_negative = (
        float(np.nanmean(purities[negative_mask]))
        if negative_mask.any()
        else np.nan
    )

    return {
        "folder_name": folder_name,
        "beta": beta,
        "dim_z": dim_z,
        "latent_path": path.as_posix(),
        "k": k,
        "distance": distance_name,
        "num_positive": int(positive_mask.sum()),
        "num_negative": int(negative_mask.sum()),
        "num_labeled": int(len(labels)),
        "avg_purity_positive": avg_purity_positive,
        "avg_purity_negative": avg_purity_negative,
        "avg_purity": float(np.nanmean(purities)),
    }


def highlight_max_purity_cells(path: Path) -> None:
    """Paint the maximum value in each purity column yellow.

    If several rows share the same maximum, all tied cells are highlighted.
    """
    workbook = load_workbook(path)
    sheet = workbook.active
    headers = [cell.value for cell in sheet[1]]
    yellow_fill = PatternFill(fill_type="solid", fgColor="FFFF00")

    for column_name in PURITY_COLUMNS:
        if column_name not in headers:
            continue

        column_index = headers.index(column_name) + 1
        values = [
            (row, sheet.cell(row=row, column=column_index).value)
            for row in range(2, sheet.max_row + 1)
        ]
        numeric_values = [
            (row, value)
            for row, value in values
            if isinstance(value, (int, float))
        ]
        if not numeric_values:
            continue

        max_value = max(value for _, value in numeric_values)
        for row, value in numeric_values:
            if value == max_value:
                sheet.cell(row=row, column=column_index).fill = yellow_fill

    workbook.save(path)


def heatmap_path(
    output_dir: Path, k: int, metric: str, value_column: str
) -> Path:
    """Build the standard output filename for a heatmap."""
    filename = f"knn_purity_k{k:02d}_{metric}_{value_column}.pdf"
    return output_dir / "figs" / filename


def make_heatmap(
    summary: pd.DataFrame,
    output_dir: Path,
    k: int,
    metric: str,
    value_column: str,
) -> None:
    """Save a beta-by-dim_z heatmap for one purity column.

    Duplicate (beta, dim_z) entries are averaged. Missing combinations remain
    NaN and are rendered as white cells.
    """
    cache_dir = output_dir / ".cache"
    cache_dir.mkdir(parents=True, exist_ok=True)
    os.environ.setdefault("MPLCONFIGDIR", str(cache_dir / "matplotlib"))
    os.environ.setdefault("XDG_CACHE_HOME", str(cache_dir))

    grid = summary.pivot_table(
        index="dim_z",
        columns="beta",
        values=value_column,
        aggfunc="mean",
    ).sort_index().sort_index(axis=1)
    if grid.empty:
        print(
            f"Skipped heatmap for {metric} {value_column}: "
            "no beta/dim_z grid could be built."
        )
        return

    output = heatmap_path(output_dir, k, metric, value_column)
    output.parent.mkdir(parents=True, exist_ok=True)

    masked_values = np.ma.masked_invalid(grid.to_numpy(dtype=float))
    cmap = plt.colormaps["viridis"].copy()
    cmap.set_bad("white")

    width = max(6.0, 1.0 + 1.0 * len(grid.columns))
    height = max(4.0, 1.0 + 0.45 * len(grid.index))
    fig, ax = plt.subplots(figsize=(width, height))
    image = ax.imshow(masked_values, cmap=cmap, vmin=0.0, vmax=1.0)

    # ax.set_title(f"{metric} - {value_column} - k={k}")
    ax.set_title(f"Average $k$-NN purity")
    ax.set_xlabel(r"$\beta$")
    ax.set_ylabel(r"$d_z$")
    ax.set_xticks(np.arange(len(grid.columns)))
    ax.set_xticklabels([f"{value:g}" for value in grid.columns])
    ax.set_yticks(np.arange(len(grid.index)))
    ax.set_yticklabels([str(value) for value in grid.index])

    for row_index, dim_z in enumerate(grid.index):
        for col_index, beta in enumerate(grid.columns):
            value = grid.loc[dim_z, beta]
            if pd.notna(value):
                ax.text(
                    col_index,
                    row_index,
                    f"{value:.3f}",
                    ha="center",
                    va="center",
                    color="white" if value < 0.75 else "black",
                    fontsize=8,
                )

    fig.colorbar(image, ax=ax)   # , label=value_column)
    fig.tight_layout()
    fig.savefig(output, dpi=800)
    plt.close(fig)


def make_metric_heatmaps(
    summary: pd.DataFrame, output_dir: Path, k: int, metric: str
) -> None:
    """Save all heatmaps requested for a metric summary."""
    make_heatmap(summary, output_dir, k, metric, "avg_purity")
    make_heatmap(summary, output_dir, k, metric, "avg_purity_positive")


def output_path(output_dir: Path, k: int, name: str) -> Path:
    """Build the standard output filename for a kNN result table."""
    return output_dir / f"knn_purity_k{k:02d}_{name}.xlsx"


def best_purity_rows(
    summaries: list[tuple[Path, pd.DataFrame]], purity_column: str
) -> pd.DataFrame:
    """Extract the row with the highest selected purity from each summary."""
    rows = []
    for path, summary in summaries:
        max_purity = summary[purity_column].max()
        best_rows = summary[summary[purity_column].eq(max_purity)].copy()
        best_rows.insert(0, "summary_file", path.name)
        rows.append(best_rows)

    if not rows:
        return pd.DataFrame()
    return pd.concat(rows, ignore_index=True)


def main() -> None:
    args = parse_args()
    if args.k < 1:
        raise ValueError("k must be at least 1")

    paths = data_paths(args.root)
    if not paths:
        raise FileNotFoundError(
            f"No latent.parquet or vae.parquet files found under {args.root}"
        )

    args.output_dir.mkdir(parents=True, exist_ok=True)
    summaries = []
    for distance_name in args.distances:
        # Reset the seed so Monte Carlo output is reproducible for each metric
        # and independent of which other distances are requested.
        rng = np.random.default_rng(args.seed)
        rows = []
        for path in paths:
            try:
                rows.append(
                    summarize_file(
                        path,
                        args.root,
                        args.annotations,
                        args.k,
                        distance_name,
                        args.js_samples,
                        rng,
                    )
                )
            except ValueError as error:
                print(f"Skipped {path} for {distance_name}: {error}")
        if not rows:
            print(f"Skipped {distance_name}: no valid rows.")
            continue
        summary = (
            pd.DataFrame(rows)
            .sort_values("folder_name")
            .reset_index(drop=True)
        )
        output = output_path(args.output_dir, args.k, distance_name)
        summary.to_excel(output, index=False)
        highlight_max_purity_cells(output)
        make_metric_heatmaps(summary, args.output_dir, args.k, distance_name)
        summaries.append((output, summary))
        print(f"Wrote {len(summary)} rows to {output}")

    best_tables = {
        "best_avg_purity": "avg_purity",
        "best_avg_purity_positive": "avg_purity_positive",
    }
    for output_name, purity_column in best_tables.items():
        best_rows = best_purity_rows(summaries, purity_column)
        best_output = output_path(args.output_dir, args.k, output_name)
        best_rows.to_excel(best_output, index=False)
        print(f"Wrote {len(best_rows)} rows to {best_output}")


if __name__ == "__main__":
    main()
